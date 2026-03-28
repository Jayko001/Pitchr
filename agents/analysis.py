"""Analysis Agent: prompts Claude to design an Excel comps model, writes it with openpyxl."""
from __future__ import annotations

import json
import os
from typing import Optional, Callable

import anthropic
import openpyxl
from openpyxl.styles import Font, PatternFill

from agents.models import CompRecord, ValuationRange, ExcelInstructions, CellWrite

SYSTEM_PROMPT = """You are a senior investment banker specializing in early-stage startup valuations.
Given a list of comparable companies and a startup description, return JSON instructions
for building a 3-tab Excel model: Comps, Multiples, and Valuation.

Return ONLY valid JSON matching this schema. No prose, no markdown fences, just JSON:
{
  "sheets": ["Comps", "Multiples", "Valuation"],
  "writes": [
    {"sheet": "Comps", "cell": "A1", "value": "Company", "bold": true},
    ...
  ],
  "column_widths": {"Comps": {"A": 25.0, "B": 15.0}, ...}
}

Rules:
- Tab 1 (Comps): header row + one row per comp. Columns: Company, Valuation ($M), ARR ($M), EV/ARR, Last Round, Lead Investors
- Tab 2 (Multiples): median/mean/high/low for EV/ARR and EV/EBITDA
- Tab 3 (Valuation): Low/Mid/High implied valuation for the target startup using median multiple x target ARR
- Bold header rows. Use bg_color "4472C4" for headers.
- Format valuation cells with number_format "$#,##0"
- Format multiple cells with number_format "#,##0.0x"
- The Valuation sheet MUST have rows labeled exactly "Low", "Mid", "High" in column A with numeric values in column B
"""


def call_claude_for_excel_instructions(
    comps: list[CompRecord],
    startup_description: str,
) -> ExcelInstructions:
    client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
    comps_json = json.dumps([c.model_dump() for c in comps], indent=2)
    user_msg = f"""Comparable companies:
{comps_json}

Target startup description:
{startup_description}

Return the ExcelInstructions JSON."""

    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=4096,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_msg}],
    )

    raw = message.content[0].text.strip()
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    return ExcelInstructions.model_validate_json(raw.strip())


def build_excel_model(instructions: ExcelInstructions, output_path: str) -> None:
    """Execute ExcelInstructions to write an .xlsx file."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name in instructions.sheets:
        wb.create_sheet(sheet_name)

    for write in instructions.writes:
        ws = wb[write.sheet]
        cell = ws[write.cell]

        if write.formula:
            cell.value = write.formula
        elif write.value is not None:
            cell.value = write.value

        font_color = "FFFFFF" if write.bg_color else "000000"
        cell.font = Font(bold=write.bold, color=font_color)

        if write.bg_color:
            cell.fill = PatternFill(
                start_color=write.bg_color,
                end_color=write.bg_color,
                fill_type="solid",
            )

        if write.number_format:
            cell.number_format = write.number_format

    for sheet_name, cols in instructions.column_widths.items():
        ws = wb[sheet_name]
        for col_letter, width in cols.items():
            ws.column_dimensions[col_letter].width = width

    wb.save(output_path)


def _extract_valuation_range(instructions: ExcelInstructions) -> ValuationRange:
    """Read Low/Mid/High rows from the Valuation sheet writes."""
    val_writes = {w.cell: w for w in instructions.writes if w.sheet == "Valuation"}
    low = mid = high = 0.0
    for row_idx in range(2, 20):
        label_write = val_writes.get(f"A{row_idx}")
        val_write = val_writes.get(f"B{row_idx}")
        if not label_write or not val_write:
            continue
        label = str(label_write.value or "").strip().lower()
        val = val_write.value
        if isinstance(val, (int, float)):
            if label == "low":
                low = float(val)
            elif label == "mid":
                mid = float(val)
            elif label == "high":
                high = float(val)
    return ValuationRange(low=low, mid=mid, high=high, methodology="EV/ARR")


def run_analysis_agent(
    comps: list[CompRecord],
    startup_description: str,
    output_path: str,
    status_callback: Optional[Callable[[str], None]] = None,
) -> tuple[ValuationRange, str]:
    """Top-level entry point called by the orchestrator. Returns (valuation_range, output_path)."""
    if status_callback:
        status_callback("Analysis Agent: asking Claude to design Excel model...")

    instructions = call_claude_for_excel_instructions(comps, startup_description)

    if status_callback:
        status_callback("Analysis Agent: writing Excel file...")

    build_excel_model(instructions, output_path)
    valuation_range = _extract_valuation_range(instructions)

    if status_callback:
        status_callback(f"Analysis Agent: done. Valuation range: ${valuation_range.low/1e6:.0f}M–${valuation_range.high/1e6:.0f}M")

    return valuation_range, output_path
